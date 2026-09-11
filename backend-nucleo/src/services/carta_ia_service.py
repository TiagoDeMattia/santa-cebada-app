"""
Servicio de base de conocimiento para el Asistente IA de carta.
Soporta: txt, md, pdf, docx, xlsx, csv.
Chat via Groq (free tier).
"""
import io
import logging
import os
from functools import lru_cache
from pathlib import Path
from typing import Dict, List, Optional

import requests
from sqlalchemy import create_engine, text
from sqlalchemy.pool import NullPool

from ..config.settings import settings

logger = logging.getLogger(__name__)

# DB en PM2_HOME (.pm2/) para evitar ProtectHome=read-only del servicio systemd.
# ReadWritePaths del systemd unit incluye /home/tiago/.pm2 pero no /home/tiago/recetario-nucleo.
_PM2_HOME = Path(os.environ.get("PM2_HOME", Path.home() / ".pm2"))
_DB_PATH = _PM2_HOME / "carta_ia.db"

@lru_cache(maxsize=1)
def _get_engine():
    return create_engine(
        f"sqlite:///{_DB_PATH.as_posix()}",
        future=True,
        connect_args={"check_same_thread": False},
        poolclass=NullPool,
    )

def get_engine():
    return _get_engine()

# ─── Schema ──────────────────────────────────────────────────────────────────

def ensure_schema():
    with get_engine().connect() as conn:
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS carta_ia_docs (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                titulo      TEXT NOT NULL,
                tipo        TEXT NOT NULL DEFAULT 'general',
                filename    TEXT,
                mime_type   TEXT,
                contenido   TEXT NOT NULL DEFAULT '',
                activo      INTEGER NOT NULL DEFAULT 1,
                orden       INTEGER NOT NULL DEFAULT 0,
                updated_at  TEXT NOT NULL DEFAULT (strftime('%Y-%m-%d %H:%M:%S', 'now')),
                updated_by  TEXT
            )
        """))
        conn.commit()


# ─── Extracción de texto ──────────────────────────────────────────────────────

SUPPORTED_EXTENSIONS = {'txt', 'md', 'text', 'pdf', 'docx', 'xlsx', 'xls', 'csv'}


def extract_text(content: bytes, filename: str) -> str:
    ext = filename.lower().rsplit('.', 1)[-1] if '.' in filename else 'txt'

    if ext in ('txt', 'md', 'text'):
        return content.decode('utf-8', errors='replace')

    if ext == 'pdf':
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            pages = [p.extract_text() or '' for p in reader.pages]
            return '\n\n'.join(p for p in pages if p.strip())
        except ImportError:
            raise ValueError("Falta pypdf. Instalar: pip install pypdf")

    if ext == 'docx':
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            raise ValueError("Falta python-docx. Instalar: pip install python-docx")

    if ext in ('xlsx', 'xls'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"=== Hoja: {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    vals = [str(v) if v is not None else '' for v in row]
                    if any(v.strip() for v in vals):
                        lines.append('\t'.join(vals))
            return '\n'.join(lines)
        except ImportError:
            raise ValueError("Falta openpyxl. Instalar: pip install openpyxl")

    if ext == 'csv':
        import csv
        decoded = content.decode('utf-8', errors='replace')
        reader = csv.reader(io.StringIO(decoded))
        return '\n'.join('\t'.join(row) for row in reader)

    # Fallback: intentar como texto
    try:
        return content.decode('utf-8', errors='replace')
    except Exception:
        raise ValueError(f"Tipo de archivo no soportado: .{ext}")


# ─── CRUD ────────────────────────────────────────────────────────────────────

def _row_to_dict(row) -> Dict:
    return dict(row._mapping) if hasattr(row, '_mapping') else dict(row)


def get_docs() -> List[Dict]:
    ensure_schema()
    with get_engine().connect() as conn:
        rows = conn.execute(text("""
            SELECT id, titulo, tipo, filename, mime_type, activo, orden,
                   updated_at, updated_by, LENGTH(contenido) AS chars
            FROM carta_ia_docs ORDER BY orden, id
        """)).fetchall()
    return [_row_to_dict(r) for r in rows]


def get_doc(doc_id: int) -> Optional[Dict]:
    ensure_schema()
    with get_engine().connect() as conn:
        row = conn.execute(
            text("SELECT * FROM carta_ia_docs WHERE id = :id"), {"id": doc_id}
        ).fetchone()
    return _row_to_dict(row) if row else None


def create_doc(
    titulo: str,
    tipo: str,
    contenido: str,
    filename: Optional[str] = None,
    mime_type: Optional[str] = None,
    updated_by: Optional[str] = None,
) -> Dict:
    ensure_schema()
    with get_engine().connect() as conn:
        result = conn.execute(text("""
            INSERT INTO carta_ia_docs
                (titulo, tipo, filename, mime_type, contenido, activo, orden, updated_at, updated_by)
            VALUES
                (:titulo, :tipo, :filename, :mime_type, :contenido, 1, 0,
                 strftime('%Y-%m-%d %H:%M:%S', 'now'), :updated_by)
        """), {
            "titulo": titulo, "tipo": tipo, "filename": filename,
            "mime_type": mime_type, "contenido": contenido, "updated_by": updated_by,
        })
        doc_id = result.lastrowid
        conn.commit()
    return get_doc(doc_id)


def update_doc(
    doc_id: int,
    titulo: Optional[str] = None,
    tipo: Optional[str] = None,
    contenido: Optional[str] = None,
    updated_by: Optional[str] = None,
) -> Optional[Dict]:
    ensure_schema()
    sets, params = [], {"id": doc_id, "updated_by": updated_by}
    if titulo is not None:
        sets.append("titulo = :titulo"); params["titulo"] = titulo
    if tipo is not None:
        sets.append("tipo = :tipo"); params["tipo"] = tipo
    if contenido is not None:
        sets.append("contenido = :contenido"); params["contenido"] = contenido
    sets += ["updated_at = strftime('%Y-%m-%d %H:%M:%S', 'now')", "updated_by = :updated_by"]
    with get_engine().connect() as conn:
        conn.execute(text(f"UPDATE carta_ia_docs SET {', '.join(sets)} WHERE id = :id"), params)
        conn.commit()
    return get_doc(doc_id)


def replace_file(
    doc_id: int,
    content: bytes,
    filename: str,
    mime_type: str,
    updated_by: Optional[str] = None,
) -> Optional[Dict]:
    texto = extract_text(content, filename)
    ensure_schema()
    with get_engine().connect() as conn:
        conn.execute(text("""
            UPDATE carta_ia_docs
            SET contenido = :contenido, filename = :filename, mime_type = :mime_type,
                updated_at = strftime('%Y-%m-%d %H:%M:%S', 'now'), updated_by = :updated_by
            WHERE id = :id
        """), {"contenido": texto, "filename": filename, "mime_type": mime_type,
               "updated_by": updated_by, "id": doc_id})
        conn.commit()
    return get_doc(doc_id)


def toggle_active(doc_id: int) -> Optional[Dict]:
    ensure_schema()
    with get_engine().connect() as conn:
        conn.execute(
            text("UPDATE carta_ia_docs SET activo = 1 - activo WHERE id = :id"), {"id": doc_id}
        )
        conn.commit()
    return get_doc(doc_id)


def delete_doc(doc_id: int) -> bool:
    ensure_schema()
    with get_engine().connect() as conn:
        result = conn.execute(
            text("DELETE FROM carta_ia_docs WHERE id = :id"), {"id": doc_id}
        )
        conn.commit()
    return result.rowcount > 0


# ─── Chat ────────────────────────────────────────────────────────────────────

_GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
_MODEL    = "llama-3.3-70b-versatile"

_SYSTEM = """Sos un asistente interno de Santa Cebada. Respondé preguntas usando la información provista abajo (documentos + estado actual de barriles).

REGLAS:
1. Usá solo la información provista. No inventes ni uses conocimiento propio.
2. Si la respuesta no está en los datos, decí: "Esa información no está disponible."
3. Respondé en español, de forma clara y directa.
4. Si te preguntan sobre el local o la carta, usá los documentos. Si preguntan sobre barriles, birras o stock actual, usá la sección "ESTADO DE BARRILES".
5. Si te preguntan algo completamente ajeno al local, respondé: "Solo puedo responder preguntas sobre la carta y el local."

{context}"""


def _build_barriles_context() -> str:
    """Arma un bloque de texto con el estado actual de barriles para el contexto de la IA."""
    try:
        from .barriles_service import get_barriles, get_info_birras_con_precio
        barriles = get_barriles(filtro="activos", sucursal_id="1")
    except Exception as e:
        logger.warning(f"No se pudo obtener barriles para contexto IA: {e}")
        return ""

    if not barriles:
        return ""

    pinchadas = [b for b in barriles if b.get("estado") == "Pinchada"]
    en_camara = [b for b in barriles if b.get("estado") == "En Camara"]

    lines = ["--- ESTADO DE BARRILES (tiempo real) ---"]

    if pinchadas:
        lines.append("\nEN CANILLA ahora mismo:")
        for b in sorted(pinchadas, key=lambda x: int(x.get("canilla") or 999)):
            canilla = b.get("canilla", "?")
            estilo  = b.get("estilo", "?")
            litros  = b.get("litros_barril", "")
            fecha   = b.get("fecha_pinchado", "")
            dias    = b.get("dias_pinchado", "")
            info = f"  Canilla {canilla}: {estilo} ({litros}L)"
            if fecha:
                info += f" — pinchada el {fecha}"
            if dias:
                info += f" (hace {dias} días)"
            lines.append(info)

    if en_camara:
        lines.append("\nEN CÁMARA (próximas a pinchar):")
        for b in en_camara:
            estilo = b.get("estilo", "?")
            litros = b.get("litros_barril", "")
            prov   = b.get("proveedor", "")
            info = f"  {estilo} ({litros}L)"
            if prov:
                info += f" — {prov}"
            lines.append(info)

    # Agregar precios si están disponibles
    try:
        precio_data = get_info_birras_con_precio(sucursal_id="1")
        birras_precio = {b["estilo"].lower(): b["precio"] for b in precio_data.get("birras", []) if b.get("precio")}
        hora_santa = precio_data.get("hora_santa_precio", "")
        if birras_precio:
            lines.append("\nPRECIOS DE BIRRAS:")
            for estilo, precio in sorted(birras_precio.items()):
                lines.append(f"  {estilo}: ${precio}")
        if hora_santa:
            lines.append(f"\nHora Santa (17-20hs): ${hora_santa} (cervezas tipo B)")
    except Exception as e:
        logger.warning(f"No se pudo obtener precios de birras: {e}")

    return "\n".join(lines)


def chat(pregunta: str) -> str:
    api_key = getattr(settings, 'groq_api_key', '')
    if not api_key:
        return "Error: GROQ_API_KEY no configurada. El administrador debe agregar la API key al servidor."

    ensure_schema()
    with get_engine().connect() as conn:
        rows = conn.execute(text(
            "SELECT titulo, tipo, contenido FROM carta_ia_docs WHERE activo = 1 ORDER BY orden, id"
        )).fetchall()

    sections = []

    if rows:
        docs_text = "\n\n".join(f"--- {r[0]} ({r[1]}) ---\n{r[2]}" for r in rows)
        sections.append("DOCUMENTOS:\n" + docs_text)
    else:
        sections.append("DOCUMENTOS:\n(Sin documentos cargados aún)")

    barriles_ctx = _build_barriles_context()
    if barriles_ctx:
        sections.append(barriles_ctx)

    if not rows and not barriles_ctx:
        return "No hay información disponible aún. Un administrador debe subir los archivos de la carta."

    context = "\n\n".join(sections)
    system_msg = _SYSTEM.format(context=context)

    try:
        resp = requests.post(
            _GROQ_URL,
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json={
                "model": _MODEL,
                "messages": [
                    {"role": "system", "content": system_msg},
                    {"role": "user",   "content": pregunta},
                ],
                "temperature": 0.1,
                "max_tokens": 1024,
            },
            timeout=30,
        )
        resp.raise_for_status()
        return resp.json()["choices"][0]["message"]["content"]
    except requests.exceptions.HTTPError as e:
        if e.response is not None and e.response.status_code == 401:
            return "Error: API key de Groq inválida. Verificá la clave en la configuración del servidor."
        if e.response is not None and e.response.status_code == 429:
            return "El servicio de IA está recibiendo muchas consultas. Esperá unos segundos y volvé a intentar."
        logger.error(f"Groq HTTP error: {e}")
        return "Error al contactar el servicio de IA. Intentá de nuevo en un momento."
    except requests.exceptions.Timeout:
        return "El servicio de IA tardó demasiado en responder. Intentá de nuevo."
    except Exception as e:
        logger.error(f"Groq error: {e}")
        return "Error inesperado al procesar tu pregunta."
